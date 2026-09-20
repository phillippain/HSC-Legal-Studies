#!/usr/bin/env python3
"""Pull the text out of the finished summaries so they can be turned into records.

Only the curated ones: the case cards, evidence summary tables and evidence sheets
that were written to be read, not the raw judgments behind them.
"""
import json, os, re, subprocess, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
lib = json.load(open(f"{ROOT}/data/library.json"))
LIB = lib["root"]
OUT = os.path.expanduser("~/work/summaries")
os.makedirs(OUT, exist_ok=True)

CURATED = re.compile(r"case summary|case card|case study card|evidence summary|"
                     r"evidence sheet|evidence file|evidence and model|comparison sheet|"
                     r"evidence bank", re.I)


def wanted(f):
    if not f.get("primary") or f["bucket"] != "evidence":
        return False
    p = f["path"]
    if p.startswith("Jade Cases/Completed Case Summaries/"):
        return True
    if p.startswith("Cases/BOCSAR/"):
        return True
    if CURATED.search(f["name"]):
        return True
    return False


def text_of(path):
    full = os.path.join(LIB, path)
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        try:
            r = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", full, "-"],
                               capture_output=True, timeout=90)
            if r.returncode == 0 and r.stdout.strip():
                return r.stdout.decode("utf-8", "replace")
        except Exception:
            pass
        import pdfplumber
        with pdfplumber.open(full) as pdf:
            return "\n".join((pg.extract_text() or "") for pg in pdf.pages[:12])
    if ext == ".docx":
        import docx
        d = docx.Document(full)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if ext == ".rtf":
        r = subprocess.run(["textutil", "-convert", "txt", "-stdout", full],
                           capture_output=True, timeout=90)
        return r.stdout.decode("utf-8", "replace") if r.returncode == 0 else ""
    return ""


def tidy(t):
    t = t.replace("\x0c", "\n")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


if __name__ == "__main__":
    picks = [f for f in lib["files"] if wanted(f)]
    picks.sort(key=lambda f: f["path"])
    manifest = []
    for i, f in enumerate(picks):
        slug = re.sub(r"[^a-z0-9]+", "-", f["name"].lower()).strip("-")[:60]
        dest = os.path.join(OUT, f"{i:03d}_{slug}.txt")
        try:
            t = tidy(text_of(f["path"]))
        except Exception as e:
            t = ""
            print(f"  !! {f['path']}: {e}", file=sys.stderr)
        open(dest, "w", encoding="utf-8").write(t)
        manifest.append({"i": i, "path": f["path"], "name": f["name"],
                         "tags": f["tags"], "kind": f["kind"],
                         "txt": dest, "chars": len(t)})
    json.dump(manifest, open(os.path.join(OUT, "manifest.json"), "w"), indent=1)
    print(f"{len(picks)} curated summaries extracted to {OUT}")
    empty = [m for m in manifest if m["chars"] < 200]
    print(f"  short or empty: {len(empty)}")
    for m in empty:
        print("    " + m["path"])
